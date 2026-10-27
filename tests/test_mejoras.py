"""
tests/test_mejoras.py – Pruebas de las dos mejoras solicitadas.

- Exportación de productos con bajo stock a PDF y Excel.
- Login/validación de usernames sin distinguir mayúsculas/minúsculas.

Todas las pruebas usan una base SQLite temporal y archivos en un directorio
temporal. No leen ni escriben la base de datos real minimarket.db.

Ejecutar:
    .\\venv\\Scripts\\python.exe tests\\test_mejoras.py
"""
from __future__ import annotations

import base64
import os
import re
import sys
import tempfile
import unittest
import zlib
from decimal import Decimal
from pathlib import Path
from unittest.mock import patch

PROJECT_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(PROJECT_ROOT))
os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

from sqlalchemy import create_engine, func
from sqlalchemy.orm import sessionmaker

from database import Base
import models  # noqa: F401  (registra todas las tablas)

from models.categoria import Categoria
from models.producto import Producto
from models.usuario import Usuario

from services.export_service import (
    exportar_productos_bajo_stock_excel,
    exportar_productos_bajo_stock_pdf,
)
from services.inventario_service import productos_con_stock_bajo
from services.usuario_service import (
    UsernameAmbiguoError,
    autenticar_usuario,
    detectar_duplicados_username,
    usuario_unico_por_username,
    username_disponible,
)

from PyQt5.QtWidgets import QApplication, QFileDialog, QMessageBox, QPushButton

_app = QApplication.instance() or QApplication([])


# ── Utilidades de prueba ─────────────────────────────────────────────────────

def extraer_texto_pdf(data: bytes) -> str:
    """Extrae texto de los streams ReportLab (ASCII85 + FlateDecode)."""
    fragmentos = []
    for match in re.finditer(rb"stream\r?\n(.*?)endstream", data, re.S):
        contenido = match.group(1).strip()
        try:
            raw = base64.a85decode(contenido, adobe=True)
            fragmentos.append(zlib.decompress(raw).decode("latin-1", errors="replace"))
        except Exception:
            continue
    return "\n".join(fragmentos)


class TempDBTestCase(unittest.TestCase):
    """Caso base con una base SQLite aislada por prueba."""

    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.tmp_path = Path(self.tmp.name)
        self.engine = create_engine(
            f"sqlite:///{self.tmp_path / 'pruebas.db'}",
            connect_args={"check_same_thread": False},
        )
        Base.metadata.create_all(self.engine)
        self.Session = sessionmaker(
            bind=self.engine,
            autoflush=False,
            autocommit=False,
            expire_on_commit=False,
        )
        self.session = self.Session()

    def tearDown(self):
        self.session.close()
        self.engine.dispose()
        self.tmp.cleanup()

    def crear_producto(self, categoria, codigo, descripcion, stock, minimo,
                       unidad="UN", activo=True) -> Producto:
        p = Producto(
            codigo_barras=codigo,
            descripcion=descripcion,
            unidad_medida=unidad,
            categoria_id=categoria.id,
            activo=activo,
        )
        p.stock_actual = Decimal(stock)
        p.stock_minimo = Decimal(minimo)
        self.session.add(p)
        return p

    def crear_usuario(self, username, password, nombre="Usuario", rol="cajero",
                      activo=True) -> Usuario:
        u = Usuario(
            username=username,
            nombre_completo=nombre,
            rol=rol,
            activo=activo,
        )
        u.set_password(password)
        self.session.add(u)
        return u


# ── Mejora 1: exportación de bajo stock ──────────────────────────────────────

class ExportacionBajoStockTest(TempDBTestCase):

    def setUp(self):
        super().setUp()
        cat = Categoria(nombre="Bebidas & Jugos", descripcion="Pruebas")
        self.session.add(cat)
        self.session.flush()

        self.p1 = self.crear_producto(cat, "P001", "Agua <600ml>", "0", "10")
        self.p2 = self.crear_producto(cat, "P002", "Arroz a granel", "1.250", "2.500", unidad="KG")
        self.p3 = self.crear_producto(cat, "P003", "Producto justo", "5", "5")
        self.p4 = self.crear_producto(cat, "P004", "Producto óptimo", "20", "5")
        self.p5 = self.crear_producto(cat, "P005", "Inactivo bajo", "0", "3", activo=False)
        self.session.commit()

        self.bajos = productos_con_stock_bajo(self.session)

    def snapshot_stock(self):
        return {
            p.codigo_barras: (
                p._stock_actual_m, p._stock_minimo_m,
                p.stock_actual, p.stock_minimo,
            )
            for p in self.session.query(Producto).order_by(Producto.codigo_barras).all()
        }

    def test_filtro_solo_stock_menor_o_igual_al_minimo(self):
        codigos = [p.codigo_barras for p in self.bajos]
        self.assertEqual(codigos, ["P001", "P002", "P003"])
        self.assertNotIn("P004", codigos)  # stock 20 > mínimo 5
        self.assertNotIn("P005", codigos)  # inactivo

        faltantes = {p.codigo_barras: p.stock_minimo - p.stock_actual for p in self.bajos}
        self.assertEqual(faltantes["P001"], Decimal("10"))
        self.assertEqual(faltantes["P002"], Decimal("1.250"))
        self.assertEqual(faltantes["P003"], Decimal("0"))

    def test_exportacion_excel_correcta(self):
        import openpyxl

        antes = self.snapshot_stock()
        destino = self.tmp_path / "bajo_stock.xlsx"
        resultado = exportar_productos_bajo_stock_excel(self.bajos, str(destino))

        self.assertEqual(resultado, str(destino))
        self.assertTrue(destino.is_file())
        self.assertGreater(destino.stat().st_size, 1000)

        wb = openpyxl.load_workbook(destino)
        ws = wb.active
        self.assertEqual(ws.title, "Bajo Stock")
        self.assertIn("REPORTE DE PRODUCTOS CON BAJO STOCK", ws["A1"].value)
        self.assertIn("Generado:", ws["A2"].value)

        headers = [ws.cell(row=6, column=c).value for c in range(1, 8)]
        self.assertEqual(headers, [
            "Código", "Producto", "Categoría", "Unidad de Medida",
            "Stock Actual", "Stock Mínimo", "Faltante para Mínimo",
        ])

        # Solo 3 productos: filas 7, 8 y 9.
        self.assertEqual([ws.cell(row=r, column=1).value for r in (7, 8, 9)],
                         ["P001", "P002", "P003"])
        self.assertIsNone(ws.cell(row=10, column=1).value)

        self.assertEqual(ws.auto_filter.ref, "A6:G9")
        self.assertEqual(ws.freeze_panes, "A7")
        self.assertEqual(ws["G7"].value, 10)
        self.assertAlmostEqual(ws["E8"].value, 1.25)
        self.assertAlmostEqual(ws["F8"].value, 2.5)
        self.assertAlmostEqual(ws["G8"].value, 1.25)
        self.assertEqual(ws["G7"].number_format, "#,##0.###")
        self.assertEqual(ws["G8"].number_format, "#,##0.###")
        self.assertGreaterEqual(ws.column_dimensions["B"].width, 22)
        self.assertGreater(ws.column_dimensions["G"].width, 0)

        self.assertEqual(self.snapshot_stock(), antes, "La exportación no debe modificar stocks")

    def test_exportacion_pdf_correcta(self):
        antes = self.snapshot_stock()
        destino = self.tmp_path / "bajo_stock.pdf"
        resultado = exportar_productos_bajo_stock_pdf(self.bajos, str(destino))

        self.assertEqual(resultado, str(destino))
        self.assertTrue(destino.is_file())
        data = destino.read_bytes()
        self.assertTrue(data.startswith(b"%PDF"))
        self.assertTrue(data.rstrip().endswith(b"%%EOF"))
        self.assertGreater(len(data), 1000)

        texto = extraer_texto_pdf(data)
        self.assertIn("Reporte de Productos con Bajo Stock", texto)
        self.assertIn("Generado:", texto)
        self.assertIn("Faltante", texto)
        for codigo in ("P001", "P002", "P003"):
            self.assertIn(codigo, texto)
        self.assertNotIn("P004", texto)
        self.assertNotIn("P005", texto)

        self.assertEqual(self.snapshot_stock(), antes, "La exportación no debe modificar stocks")

    def test_sin_productos_bajo_stock(self):
        for p in self.session.query(Producto).all():
            p.stock_actual = p.stock_minimo + Decimal("5")
        self.session.commit()

        self.assertEqual(productos_con_stock_bajo(self.session), [])

        xlsx = self.tmp_path / "sin_bajos.xlsx"
        pdf = self.tmp_path / "sin_bajos.pdf"
        exportar_productos_bajo_stock_excel([], str(xlsx))
        exportar_productos_bajo_stock_pdf([], str(pdf))

        import openpyxl
        ws = openpyxl.load_workbook(xlsx).active
        self.assertIsNone(ws.auto_filter.ref)
        self.assertIn("No se encontraron productos", ws["A7"].value)

        texto = extraer_texto_pdf(pdf.read_bytes())
        self.assertIn("No se encontraron productos", texto)

    def test_ruta_de_guardado_invalida(self):
        xlsx = self.tmp_path / "carpeta_inexistente" / "reporte.xlsx"
        pdf = self.tmp_path / "carpeta_inexistente" / "reporte.pdf"

        with self.assertRaises(Exception):
            exportar_productos_bajo_stock_excel(self.bajos, str(xlsx))
        with self.assertRaises(Exception):
            exportar_productos_bajo_stock_pdf(self.bajos, str(pdf))

        self.assertFalse(xlsx.exists())
        self.assertFalse(pdf.exists())

    def test_error_de_permisos_real(self):
        """Un archivo de destino de solo lectura debe provocar PermissionError."""
        import stat

        for destino, funcion in (
            (self.tmp_path / "solo_lectura.xlsx", exportar_productos_bajo_stock_excel),
            (self.tmp_path / "solo_lectura.pdf", exportar_productos_bajo_stock_pdf),
        ):
            destino.write_bytes(b"bloqueado")
            os.chmod(destino, stat.S_IREAD)
            try:
                with self.assertRaises(PermissionError):
                    funcion(self.bajos, str(destino))
            finally:
                os.chmod(destino, stat.S_IWRITE | stat.S_IREAD)


class InventarioWidgetExportTest(TempDBTestCase):
    """Prueba el cableado real de los botones/diálogos de la pestaña Inventario."""

    def setUp(self):
        super().setUp()
        from ui.inventario_widget import InventarioWidget

        cat = Categoria(nombre="Abarrotes", descripcion="Pruebas")
        self.session.add(cat)
        self.session.flush()
        self.crear_producto(cat, "UI001", "Azúcar 1kg", "0.500", "2")
        self.crear_producto(cat, "UI002", "Fideo", "10", "3")  # no bajo
        self.usuario = self.crear_usuario("cajero1", "Clave123", rol="cajero")
        self.session.commit()

        self.patcher = patch("ui.inventario_widget.SessionLocal", self.Session)
        self.patcher.start()
        self.widget = InventarioWidget(self.usuario)

    def tearDown(self):
        self.widget.close()
        self.patcher.stop()
        super().tearDown()

    def test_botones_visibles(self):
        textos = [b.text() for b in self.widget.findChildren(QPushButton)]
        self.assertIn("📄 Exportar PDF", textos)
        self.assertIn("📊 Exportar Excel", textos)

    def test_exportar_excel_desde_ui(self):
        destino = self.tmp_path / "ui_bajo.xlsx"
        with patch("ui.inventario_widget.QFileDialog.getSaveFileName",
                   return_value=(str(destino), "")) as dialogo, \
             patch("ui.inventario_widget.QMessageBox.question",
                   return_value=QMessageBox.No) as pregunta, \
             patch("ui.inventario_widget.QMessageBox.critical") as critico:
            self.widget._exportar_bajo_stock_excel()

        dialogo.assert_called_once()
        pregunta.assert_called_once()
        critico.assert_not_called()
        self.assertTrue(destino.is_file())

    def test_exportar_pdf_desde_ui(self):
        destino = self.tmp_path / "ui_bajo.pdf"
        with patch("ui.inventario_widget.QFileDialog.getSaveFileName",
                   return_value=(str(destino), "")) as dialogo, \
             patch("ui.inventario_widget.QMessageBox.question",
                   return_value=QMessageBox.No) as pregunta, \
             patch("ui.inventario_widget.QMessageBox.critical") as critico:
            self.widget._exportar_bajo_stock_pdf()

        dialogo.assert_called_once()
        pregunta.assert_called_once()
        critico.assert_not_called()
        self.assertTrue(destino.is_file())
        self.assertTrue(destino.read_bytes().startswith(b"%PDF"))

    def test_ui_informa_si_no_hay_bajo_stock(self):
        for p in self.session.query(Producto).all():
            p.stock_actual = p.stock_minimo + Decimal("10")
        self.session.commit()
        self.widget.session.expire_all()

        with patch("ui.inventario_widget.QFileDialog.getSaveFileName") as dialogo, \
             patch("ui.inventario_widget.QMessageBox.information") as info:
            self.widget._exportar_bajo_stock_excel()

        dialogo.assert_not_called()
        info.assert_called_once()
        self.assertIn("No hay productos", info.call_args[0][2])
        self.assertIn("No se generó ningún archivo", info.call_args[0][2])

    def test_ui_informa_error_de_permisos(self):
        destino = self.tmp_path / "sin_permiso.xlsx"
        with patch("ui.inventario_widget.QFileDialog.getSaveFileName",
                   return_value=(str(destino), "")), \
             patch("ui.inventario_widget.exportar_productos_bajo_stock_excel",
                   side_effect=PermissionError("Acceso denegado")), \
             patch("ui.inventario_widget.QMessageBox.critical") as critico, \
             patch("ui.inventario_widget.QMessageBox.question") as pregunta:
            self.widget._exportar_bajo_stock_excel()

        critico.assert_called_once()
        self.assertIn("permisos", critico.call_args[0][1].lower())
        pregunta.assert_not_called()
        self.assertFalse(destino.exists())


# ── Mejora 2: autenticación sin distinguir mayúsculas/minúsculas ─────────────

class AutenticacionTest(TempDBTestCase):

    def setUp(self):
        super().setUp()
        self.admin = self.crear_usuario("admin", "Secret123", "Administrador", "admin")
        self.patrick = self.crear_usuario("patrick", "Cajero#1", "Patrick", "admin,cajero")
        self.inactivo = self.crear_usuario("inactivo", "Pass123", "Inactivo", "cajero", activo=False)
        self.session.commit()

    def test_variantes_de_mayusculas_identifican_la_misma_cuenta(self):
        for variante in ("admin", "ADMIN", "Admin", "aDmIn"):
            usuario = autenticar_usuario(self.session, variante, "Secret123")
            self.assertIsNotNone(usuario, variante)
            self.assertEqual(usuario.id, self.admin.id, variante)

    def test_password_sigue_siendo_sensible_a_mayusculas(self):
        self.assertIsNone(autenticar_usuario(self.session, "admin", "secret123"))
        self.assertIsNone(autenticar_usuario(self.session, "ADMIN", "SECRET123"))
        self.assertIsNotNone(autenticar_usuario(self.session, "ADMIN", "Secret123"))

    def test_roles_se_conservan(self):
        usuario = autenticar_usuario(self.session, "ADMIN", "Secret123")
        self.assertEqual(usuario.get_roles(), ["admin"])
        self.assertTrue(usuario.tiene_rol("admin"))

        patrick = autenticar_usuario(self.session, "PATRICK", "Cajero#1")
        self.assertEqual(patrick.get_roles(), ["admin", "cajero"])
        self.assertTrue(patrick.tiene_rol("admin"))
        self.assertTrue(patrick.tiene_rol("cajero"))

    def test_usuario_inactivo_no_autentica(self):
        self.assertIsNone(autenticar_usuario(self.session, "INACTIVO", "Pass123"))

    def test_duplicados_por_mayusculas_bloquean_login_ambiguo(self):
        self.crear_usuario("ADMIN", "OtraClave", "Duplicado", "admin")
        self.session.commit()

        duplicados = detectar_duplicados_username(self.session)
        self.assertIn("admin", duplicados)
        self.assertEqual({u.username for u in duplicados["admin"]}, {"admin", "ADMIN"})

        # Ninguna variante puede autenticarse mientras exista ambigüedad.
        for variante in ("admin", "ADMIN", "Admin"):
            with self.assertRaises(UsernameAmbiguoError):
                autenticar_usuario(self.session, variante, "Secret123")
            with self.assertRaises(UsernameAmbiguoError):
                autenticar_usuario(self.session, variante, "OtraClave")

        self.assertFalse(username_disponible(self.session, "AdMiN"))
        self.assertIsNone(usuario_unico_por_username(self.session, "Admin"))

    def test_validacion_de_creacion_y_edicion(self):
        self.assertFalse(username_disponible(self.session, "admin"))
        self.assertFalse(username_disponible(self.session, "ADMIN"))
        self.assertFalse(username_disponible(self.session, "  AdMiN  "))
        self.assertTrue(username_disponible(self.session, "admin", excluir_id=self.admin.id))
        self.assertTrue(username_disponible(self.session, "nuevo_usuario"))
        # Editar a patrick con el nombre de admin debe rechazarse.
        self.assertFalse(username_disponible(self.session, "ADMIN", excluir_id=self.patrick.id))


class LoginDialogTest(TempDBTestCase):

    def setUp(self):
        super().setUp()
        self.admin = self.crear_usuario("admin", "Secret123", "Administrador", "admin")
        self.session.commit()
        self.patcher = patch("ui.login_dialog.SessionLocal", self.Session)
        self.patcher.start()

    def tearDown(self):
        self.patcher.stop()
        super().tearDown()

    def _login(self, username, password):
        from ui.login_dialog import LoginDialog

        dlg = LoginDialog()
        dlg._shake = lambda: None
        dlg.input_user.setText(username)
        dlg.input_pass.setText(password)
        dlg._on_login()
        usuario = dlg.usuario_autenticado
        resultado = dlg.result()
        dlg.close()
        return usuario, resultado

    def test_login_variantes(self):
        from PyQt5.QtWidgets import QDialog

        for variante in ("admin", "ADMIN", "AdMiN"):
            usuario, resultado = self._login(variante, "Secret123")
            self.assertIsNotNone(usuario, variante)
            self.assertEqual(usuario.id, self.admin.id)
            self.assertEqual(resultado, QDialog.Accepted)

    def test_login_password_incorrecta(self):
        from PyQt5.QtWidgets import QDialog

        usuario, resultado = self._login("ADMIN", "secret123")
        self.assertIsNone(usuario)
        self.assertNotEqual(resultado, QDialog.Accepted)

    def test_login_ambiguo_no_autentica(self):
        self.crear_usuario("ADMIN", "OtraClave", "Duplicado", "admin")
        self.session.commit()

        from PyQt5.QtWidgets import QDialog

        with patch("ui.login_dialog.QMessageBox.warning") as aviso:
            usuario, resultado = self._login("Admin", "Secret123")

        self.assertIsNone(usuario)
        self.assertNotEqual(resultado, QDialog.Accepted)
        aviso.assert_called_once()


class UsuarioDialogTest(TempDBTestCase):

    def setUp(self):
        super().setUp()
        self.admin = self.crear_usuario("admin", "Secret123", "Administrador", "admin")
        self.cajero = self.crear_usuario("cajero1", "Cajero#1", "Cajero Uno", "cajero")
        self.session.commit()

    def test_creacion_rechaza_username_equivalente(self):
        from ui.usuarios_widget import UsuarioDialog

        dlg = UsuarioDialog(self.session)
        dlg.input_username.setText("ADMIN")
        dlg.input_nombre.setText("Otro Admin")
        dlg.input_pass.setText("Nueva123")

        with patch("ui.usuarios_widget.QMessageBox.warning") as aviso:
            dlg._guardar()

        aviso.assert_called_once()
        self.assertEqual(
            self.session.query(Usuario).filter(func.lower(Usuario.username) == "admin").count(),
            1,
        )
        dlg.close()

    def test_edicion_rechaza_username_equivalente(self):
        from ui.usuarios_widget import UsuarioDialog

        dlg = UsuarioDialog(self.session, self.cajero)
        dlg.input_username.setText("AdMiN")

        with patch("ui.usuarios_widget.QMessageBox.warning") as aviso:
            dlg._guardar()

        aviso.assert_called_once()
        self.session.rollback()
        cajero = self.session.get(Usuario, self.cajero.id)
        self.assertEqual(cajero.username, "cajero1")
        dlg.close()

    def test_edicion_del_mismo_usuario_conserva_username_y_rol(self):
        from ui.usuarios_widget import UsuarioDialog

        dlg = UsuarioDialog(self.session, self.admin)
        dlg.input_nombre.setText("Administrador Actualizado")

        with patch("ui.usuarios_widget.QMessageBox.warning") as aviso:
            dlg._guardar()

        aviso.assert_not_called()
        self.session.expire_all()
        admin = self.session.get(Usuario, self.admin.id)
        self.assertEqual(admin.username, "admin")
        self.assertEqual(admin.rol, "admin")
        self.assertTrue(admin.verificar_password("Secret123"))
        dlg.close()


class UsuariosWidgetDuplicadosTest(TempDBTestCase):

    def setUp(self):
        super().setUp()
        self.admin = self.crear_usuario("admin", "Secret123", "Administrador", "admin")
        self.crear_usuario("ADMIN", "OtraClave", "Duplicado", "admin")
        self.session.commit()
        self.patcher = patch("ui.usuarios_widget.SessionLocal", self.Session)
        self.patcher.start()

    def tearDown(self):
        self.patcher.stop()
        super().tearDown()

    def test_panel_informa_duplicados_sin_borrar(self):
        from ui.usuarios_widget import UsuariosWidget

        with patch("ui.usuarios_widget.QMessageBox.warning") as aviso:
            widget = UsuariosWidget(self.admin)

        aviso.assert_called_once()
        self.assertIn("duplicados", aviso.call_args[0][1].lower())
        self.assertEqual(self.session.query(Usuario).count(), 2)
        widget.close()


if __name__ == "__main__":
    unittest.main(verbosity=2)
