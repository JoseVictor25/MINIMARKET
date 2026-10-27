"""
tests/test_caja_sesiones_y_actualizacion.py – Pruebas de cajas y refresco entre módulos.

Cubre:
  1. Propiedad de cajas (A no se hereda a B; venta con caja ajena rechazada).
  2. Una sola caja abierta por usuario.
  3. Historial de ventas actualizado sin cambiar filtros y sin duplicados.
  4. Inventario/alertas/POS sincronizados tras ventas.
  5. Compras y ajustes de inventario.
  6. Exportaciones PDF/Excel con stock actualizado y cantidades fraccionarias.
  7. Coordinación central de señales en MainWindow.

Todas las pruebas usan SQLite temporal. No leen ni escriben minimarket.db.
Ejecutar:
    .\\venv\\Scripts\\python.exe tests\\test_caja_sesiones_y_actualizacion.py
"""
from __future__ import annotations

import base64
import os
import re
import sys
import tempfile
import unittest
import zlib
from contextlib import contextmanager
from decimal import Decimal
from pathlib import Path
from unittest.mock import patch

PROJECT_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(PROJECT_ROOT))
os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker

from database import Base
import models  # noqa: F401  (registra todas las tablas)

from models.ajuste_inventario import AjusteInventario
from models.caja import SesionCaja
from models.categoria import Categoria
from models.compra import Compra
from models.producto import Producto
from models.proveedor import Proveedor
from models.usuario import Usuario
from models.venta import Venta

from services.caja_service import (
    CajaNoAutorizadaError,
    CajaYaAbiertaError,
    abrir_caja,
    cajas_abiertas,
    cajas_abiertas_usuario,
    cerrar_caja,
    obtener_caja_activa,
)
from services.compra_service import registrar_compra
from services.export_service import (
    exportar_productos_bajo_stock_excel,
    exportar_productos_bajo_stock_pdf,
)
from services.inventario_service import ajustar_stock, productos_con_stock_bajo
from services.venta_service import CajaNoValidaError, registrar_venta

from PyQt5.QtCore import pyqtSignal
from PyQt5.QtWidgets import QApplication, QWidget

_app = QApplication.instance() or QApplication([])


def extraer_texto_pdf(data: bytes) -> str:
    """Extrae texto de los streams ReportLab (ASCII85 + FlateDecode)."""
    fragmentos = []
    for match in re.finditer(rb"stream\r?\n(.*?)endstream", data, re.S):
        contenido = match.group(1).strip()
        try:
            raw = base64.a85decode(contenido, adobe=True)
            fragmentos.append(zlib.decompress(raw).decode("latin-1", errors="replace"))
        except Exception:
            continue
    return "\n".join(fragmentos)


class TempDBTestCase(unittest.TestCase):
    """Caso base con base SQLite aislada y sesiones de corta duración."""

    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.tmp_path = Path(self.tmp.name)
        self.engine = create_engine(
            f"sqlite:///{self.tmp_path / 'pruebas.db'}",
            connect_args={"check_same_thread": False},
        )
        Base.metadata.create_all(self.engine)
        self.Session = sessionmaker(
            bind=self.engine,
            autoflush=False,
            autocommit=False,
            expire_on_commit=False,
        )
        self.session = self.Session()

    def tearDown(self):
        self.session.close()
        self.engine.dispose()
        self.tmp.cleanup()

    # ── Utilidades ───────────────────────────────────────────────────────────
    def crear_usuario(self, username, rol="cajero", nombre=None) -> Usuario:
        u = Usuario(
            username=username,
            nombre_completo=nombre or username.title(),
            rol=rol,
            activo=True,
        )
        u.set_password("Clave123")
        self.session.add(u)
        self.session.flush()
        return u

    def crear_categoria(self, nombre="Pruebas") -> Categoria:
        c = Categoria(nombre=nombre, descripcion="Categoría de prueba")
        self.session.add(c)
        self.session.flush()
        return c

    def crear_producto(self, categoria, codigo, descripcion, stock, minimo,
                       unidad="UN", precio_compra="5.00",
                       precio_venta="10.00") -> Producto:
        p = Producto(
            codigo_barras=codigo,
            descripcion=descripcion,
            unidad_medida=unidad,
            categoria_id=categoria.id,
            activo=True,
        )
        p.stock_actual = Decimal(str(stock))
        p.stock_minimo = Decimal(str(minimo))
        p.precio_compra_bs = Decimal(str(precio_compra))
        p.precio_venta_bs = Decimal(str(precio_venta))
        self.session.add(p)
        self.session.flush()
        return p

    def abrir_caja(self, usuario, turno="Mañana") -> SesionCaja:
        caja = abrir_caja(self.session, usuario.id, turno=turno)
        self.session.commit()
        return caja

    def vender(self, usuario, producto, cantidad, caja_id):
        venta = registrar_venta(
            self.session,
            usuario,
            [{"producto": producto, "cantidad": Decimal(str(cantidad))}],
            metodo_pago="efectivo",
            sesion_caja_id=caja_id,
        )
        self.session.commit()
        return venta

    @contextmanager
    def fake_get_session(self):
        """Reemplazo de database.get_session para widgets, sobre la base temporal."""
        s = self.Session()
        try:
            yield s
            s.commit()
        except Exception:
            s.rollback()
            raise
        finally:
            s.close()


# ══════════════════════════════════════════════════════════════════════════════
#  PRUEBA 1 – PROPIEDAD DE CAJAS
# ══════════════════════════════════════════════════════════════════════════════

class PropiedadCajasTest(TempDBTestCase):

    def setUp(self):
        super().setUp()
        self.cat = self.crear_categoria()
        self.prod = self.crear_producto(self.cat, "P001", "Producto prueba", 20, 5)
        self.a = self.crear_usuario("usuario_a", "admin,cajero", "Usuario A")
        self.b = self.crear_usuario("usuario_b", "admin", "Usuario B")
        self.session.commit()
        self.caja_a = self.abrir_caja(self.a)

    def test_b_no_hereda_la_caja_de_a(self):
        self.assertIsNone(obtener_caja_activa(self.session, self.b.id))
        self.assertEqual(cajas_abiertas_usuario(self.session, self.b.id), [])
        # Supervisión: la caja de A sigue visible para consulta global.
        self.assertIn(self.caja_a.id, [c.id for c in cajas_abiertas(self.session)])

    def test_b_no_puede_vender_con_caja_ajena(self):
        with self.assertRaises(CajaNoValidaError):
            registrar_venta(
                self.session,
                self.b,
                [{"producto": self.prod, "cantidad": Decimal("1")}],
                metodo_pago="efectivo",
                sesion_caja_id=self.caja_a.id,
            )
        self.session.rollback()
        # No se creó ninguna venta ni se descontó stock.
        self.assertEqual(self.session.query(Venta).count(), 0)
        self.session.refresh(self.prod)
        self.assertEqual(self.prod.stock_actual, Decimal("20"))

    def test_venta_rechaza_caja_inexistente_o_cerrada(self):
        with self.assertRaises(CajaNoValidaError):
            registrar_venta(
                self.session, self.a,
                [{"producto": self.prod, "cantidad": Decimal("1")}],
                sesion_caja_id=999999,
            )
        self.session.rollback()

        cerrar_caja(
            self.session, self.caja_a.id, Decimal("0.00"),
            usuario_id=self.a.id,
        )
        self.session.commit()

        with self.assertRaises(CajaNoValidaError):
            registrar_venta(
                self.session, self.a,
                [{"producto": self.prod, "cantidad": Decimal("1")}],
                sesion_caja_id=self.caja_a.id,
            )
        self.session.rollback()
        self.assertEqual(self.session.query(Venta).count(), 0)

    def test_solo_el_propietario_puede_cerrar_su_caja(self):
        with self.assertRaises(CajaNoAutorizadaError):
            cerrar_caja(
                self.session, self.caja_a.id, Decimal("0.00"),
                usuario_id=self.b.id,
            )
        self.session.rollback()
        self.session.refresh(self.caja_a)
        self.assertEqual(self.caja_a.estado, SesionCaja.ESTADO_ABIERTA)

        cerrar_caja(
            self.session, self.caja_a.id, Decimal("0.00"),
            usuario_id=self.a.id,
        )
        self.session.commit()
        self.session.refresh(self.caja_a)
        self.assertEqual(self.caja_a.estado, SesionCaja.ESTADO_CERRADA)
        self.assertIsNone(obtener_caja_activa(self.session, self.a.id))

    def test_a_recupera_su_caja_y_puede_vender(self):
        # Equivale a cerrar sesión de A y volver a entrar: la caja permanece.
        self.assertEqual(obtener_caja_activa(self.session, self.a.id).id, self.caja_a.id)
        venta = self.vender(self.a, self.prod, 1, self.caja_a.id)
        self.assertEqual(venta.sesion_caja_id, self.caja_a.id)
        self.session.refresh(self.prod)
        self.assertEqual(self.prod.stock_actual, Decimal("19"))

    def test_b_puede_abrir_su_propia_caja(self):
        caja_b = self.abrir_caja(self.b)
        self.assertEqual(obtener_caja_activa(self.session, self.b.id).id, caja_b.id)
        self.assertEqual(obtener_caja_activa(self.session, self.a.id).id, self.caja_a.id)
        self.assertEqual(len(cajas_abiertas(self.session)), 2)

    def test_pos_de_b_no_toma_la_caja_de_a(self):
        from ui.pos_widget import POSWidget

        with patch("ui.pos_widget.SessionLocal", self.Session), \
             patch("ui.pos_widget.get_session", self.fake_get_session):
            pos_b = POSWidget(self.b)

        self.assertIsNone(pos_b._caja_activa_id)
        self.assertIn("Caja Cerrada", pos_b.badge_caja.text())
        pos_b.close()

    def test_tarjeta_de_caja_de_b_admin_no_muestra_la_caja_de_a(self):
        from PyQt5.QtWidgets import QLabel
        from ui.caja_widget import CajaWidget

        with patch("ui.caja_widget.SessionLocal", self.Session), \
             patch("ui.caja_widget.get_session", self.fake_get_session):
            tarjeta_b = CajaWidget(self.b)

            textos = [lbl.text() for lbl in tarjeta_b.card_actual.findChildren(QLabel)]
            self.assertTrue(any("No tienes una caja abierta" in t for t in textos))
            self.assertTrue(any("Supervisión" in t for t in textos))
            self.assertFalse(any(f"CAJA ABIERTA #{self.caja_a.id}" in t for t in textos))

            # Al abrir su propia caja, la tarjeta muestra la de B, no la de A.
            caja_b = self.abrir_caja(self.b)
            tarjeta_b.actualizar()
            textos = [lbl.text() for lbl in tarjeta_b.card_actual.findChildren(QLabel)]
            self.assertTrue(any(f"CAJA ABIERTA #{caja_b.id}" in t for t in textos))
            self.assertFalse(any(f"CAJA ABIERTA #{self.caja_a.id}" in t for t in textos))

        tarjeta_b.close()


# ══════════════════════════════════════════════════════════════════════════════
#  PRUEBA 2 – APERTURAS DUPLICADAS
# ══════════════════════════════════════════════════════════════════════════════

class AperturasDuplicadasTest(TempDBTestCase):

    def setUp(self):
        super().setUp()
        self.a = self.crear_usuario("dup_a", "cajero", "Duplicado A")
        self.b = self.crear_usuario("dup_b", "cajero", "Duplicado B")
        self.session.commit()

    def test_segunda_apertura_se_rechaza_sin_alterar_la_primera(self):
        primera = self.abrir_caja(self.a, turno="Mañana")

        with self.assertRaises(CajaYaAbiertaError) as ctx:
            abrir_caja(self.session, self.a.id, turno="Tarde")
        self.session.rollback()

        self.assertEqual([c.id for c in ctx.exception.cajas], [primera.id])
        self.session.refresh(primera)
        self.assertEqual(primera.estado, SesionCaja.ESTADO_ABIERTA)
        self.assertEqual(primera.turno, "Mañana")
        self.assertEqual(len(cajas_abiertas_usuario(self.session, self.a.id)), 1)

    def test_b_puede_abrir_su_propia_caja_sin_afectar_a_a(self):
        caja_a = self.abrir_caja(self.a)
        caja_b = self.abrir_caja(self.b, turno="Tarde")

        self.assertNotEqual(caja_a.id, caja_b.id)
        self.assertEqual(caja_b.usuario_id, self.b.id)
        self.assertEqual(caja_a.usuario_id, self.a.id)
        self.assertEqual(len(cajas_abiertas_usuario(self.session, self.a.id)), 1)
        self.assertEqual(len(cajas_abiertas_usuario(self.session, self.b.id)), 1)

    def test_error_duplicado_conserva_datos_tras_cerrar_sesion(self):
        """La UI lee el error después del context manager; no debe fallar por ORM expirado."""
        from database import get_session

        with patch("database.SessionLocal", self.Session):
            with get_session() as s:
                original = abrir_caja(s, self.a.id, turno="Mañana")

            with self.assertRaises(CajaYaAbiertaError) as ctx:
                with get_session() as s:
                    abrir_caja(s, self.a.id, turno="Tarde")

        # Acceso posterior al cierre de la sesión (como hace AperturaCajaDialog).
        self.assertEqual(ctx.exception.cajas[0].id, original.id)
        self.assertEqual(ctx.exception.cajas[0].turno, "Mañana")

    def test_ui_informa_duplicados_y_ofrece_cierre_manual(self):
        from PyQt5.QtWidgets import QLabel, QPushButton
        from ui.caja_widget import CajaWidget

        self.abrir_caja(self.a, turno="Mañana")
        # Simula un duplicado histórico preexistente (sin usar el servicio).
        dup = SesionCaja(
            usuario_id=self.a.id, turno="Tarde", fecha_apertura=None,
            estado=SesionCaja.ESTADO_ABIERTA,
        )
        dup.monto_inicial_bs = Decimal("0.00")
        self.session.add(dup)
        self.session.commit()

        with patch("ui.caja_widget.SessionLocal", self.Session), \
             patch("ui.caja_widget.get_session", self.fake_get_session):
            w = CajaWidget(self.a)

        textos = [lbl.text() for lbl in w.card_actual.findChildren(QLabel)]
        botones = [b.text() for b in w.card_actual.findChildren(QPushButton)]
        self.assertTrue(any("Se detectaron 2 cajas abiertas" in t for t in textos))
        self.assertTrue(any("cerrar cajas adicionales" in b.lower() for b in botones))
        # No se modificó ni cerró el duplicado automáticamente.
        self.session.refresh(dup)
        self.assertEqual(dup.estado, SesionCaja.ESTADO_ABIERTA)
        w.close()


# ══════════════════════════════════════════════════════════════════════════════
#  PRUEBA 3 – HISTORIAL DE VENTAS
# ══════════════════════════════════════════════════════════════════════════════

class HistorialVentasTest(TempDBTestCase):

    def setUp(self):
        super().setUp()
        self.cat = self.crear_categoria()
        self.prod = self.crear_producto(self.cat, "H001", "Producto historial", 100, 5)
        self.user = self.crear_usuario("historial", "cajero", "Cajero Historial")
        self.session.commit()
        self.caja = self.abrir_caja(self.user)

    def _nueva_venta_externa(self, cantidad=1):
        """Venta desde otra sesión, como ocurre con el POS en paralelo."""
        s = self.Session()
        try:
            u = s.get(Usuario, self.user.id)
            p = s.get(Producto, self.prod.id)
            v = registrar_venta(
                s, u, [{"producto": p, "cantidad": Decimal(str(cantidad))}],
                metodo_pago="efectivo", sesion_caja_id=self.caja.id,
            )
            s.commit()
            return v.id
        finally:
            s.close()

    def test_todos_los_registros_se_actualiza_sin_cambiar_filtro(self):
        from ui.ventas_widget import VentasWidget

        with patch("ui.ventas_widget.SessionLocal", self.Session):
            widget = VentasWidget(self.user)

        self.assertEqual(widget.combo_periodo.currentIndex(), 0)  # Todos los registros
        self.assertEqual(widget.tabla.rowCount(), 0)

        self._nueva_venta_externa()
        widget.actualizar()

        self.assertEqual(widget.combo_periodo.currentIndex(), 0)
        self.assertEqual(widget.tabla.rowCount(), 1)
        self.assertIn("1 venta", widget.lbl_kpi_cant.text())
        widget.close()

    def test_ventas_de_hoy_mantiene_filtro_y_no_duplica(self):
        from ui.ventas_widget import VentasWidget

        with patch("ui.ventas_widget.SessionLocal", self.Session):
            widget = VentasWidget(self.user)

        widget.combo_periodo.setCurrentIndex(1)  # Hoy
        self.assertEqual(widget.combo_periodo.currentIndex(), 1)

        self._nueva_venta_externa(1)
        widget.actualizar()
        self.assertEqual(widget.combo_periodo.currentIndex(), 1)
        self.assertEqual(widget.tabla.rowCount(), 1)

        self._nueva_venta_externa(2)
        widget.actualizar()
        self.assertEqual(widget.combo_periodo.currentIndex(), 1)
        self.assertEqual(widget.tabla.rowCount(), 2)

        # Cambiar a "Todos" y comprobar que siguen siendo exactamente 2.
        widget.combo_periodo.setCurrentIndex(0)
        self.assertEqual(widget.tabla.rowCount(), 2)
        self.assertEqual(self.session.query(Venta).count(), 2)
        widget.close()

    def test_doble_confirmacion_registra_una_sola_venta(self):
        from ui.pos_widget import CobrarDialog

        dlg = CobrarDialog(
            total=Decimal("10.00"),
            session=self.session,
            items=[{"producto": self.prod, "cantidad": Decimal("1")}],
            usuario=self.user,
            cliente=None,
            sesion_caja_id=self.caja.id,
        )
        dlg.metodo_pago = "efectivo"
        dlg.monto_recibido = Decimal("10.00")
        dlg.vuelto = Decimal("0.00")

        with patch("ui.pos_widget.get_session", self.fake_get_session), \
             patch("ui.pos_widget.generar_ticket_pdf", return_value=None):
            dlg._confirmar()
            dlg._confirmar()  # segundo clic / segundo Enter

        self.assertTrue(dlg.venta_confirmada)
        self.assertEqual(self.session.query(Venta).count(), 1)
        self.session.expire_all()
        self.session.refresh(self.prod)
        self.assertEqual(self.prod.stock_actual, Decimal("99"))
        dlg.close()

    def test_pos_emite_una_sola_senal_tras_venta_exitosa(self):
        from PyQt5.QtWidgets import QDialog
        from ui import pos_widget as modulo_pos

        class FakeCobrar(QDialog):
            def __init__(self, *args, **kwargs):
                super().__init__()
                self.venta_confirmada = True
                self.venta_registrada = {
                    "numero": 1, "total": Decimal("10.00"), "metodo": "efectivo",
                    "monto_recibido": Decimal("10.00"), "vuelto": Decimal("0.00"),
                    "monto_qr": Decimal("0.00"), "monto_efectivo": Decimal("10.00"),
                    "efectivo_recibido": Decimal("0.00"), "observacion": None,
                    "cliente": "Público General", "cajero": self.__class__.__name__,
                    "qr_confirmado": False, "referencia_pago": None,
                    "comprobante_pago": None,
                }
                self.pdf_path = None

            def exec(self):
                return QDialog.Accepted

        class FakeExito(QDialog):
            def __init__(self, *args, **kwargs):
                super().__init__()

            def exec(self):
                return QDialog.Accepted

        with patch("ui.pos_widget.SessionLocal", self.Session), \
             patch("ui.pos_widget.get_session", self.fake_get_session):
            pos = modulo_pos.POSWidget(self.user)

        pos._items = [{"producto": self.prod, "cantidad": Decimal("1")}]
        eventos = []
        pos.datos_actualizados.connect(eventos.append)

        with patch.object(modulo_pos, "CobrarDialog", FakeCobrar), \
             patch.object(modulo_pos, "CobroExitosoDialog", FakeExito), \
             patch("ui.pos_widget.get_session", self.fake_get_session):
            pos._abrir_cobro()

        self.assertEqual(eventos, ["venta"])
        pos.close()


# ══════════════════════════════════════════════════════════════════════════════
#  PRUEBA 4 – INVENTARIO Y STOCK MÍNIMO
# ══════════════════════════════════════════════════════════════════════════════

class InventarioStockTest(TempDBTestCase):

    def setUp(self):
        super().setUp()
        self.cat = self.crear_categoria()
        self.prod = self.crear_producto(self.cat, "S001", "Coca Cola", 10, 3)
        self.user = self.crear_usuario("stock", "admin,cajero", "Usuario Stock")
        self.session.commit()
        self.caja = self.abrir_caja(self.user)

    def _vender(self, cantidad):
        self.vender(self.user, self.prod, cantidad, self.caja.id)

    def _inventario(self):
        from ui.inventario_widget import InventarioWidget

        with patch("ui.inventario_widget.SessionLocal", self.Session):
            return InventarioWidget(self.user)

    @staticmethod
    def _fila_codigo(tabla, codigo):
        for row in range(tabla.rowCount()):
            item = tabla.item(row, 0)
            if item and item.text() == codigo:
                return row
        return -1

    def test_alerta_y_stock_se_actualizan_tras_cada_venta(self):
        self._vender(7)  # 10 - 7 = 3 (igual al mínimo → alerta)
        inv = self._inventario()

        self.assertEqual(inv.kpi_alertas.lbl_val.text(), "1 Productos")
        self.assertTrue(inv.alert_banner.isVisibleTo(inv) or not inv.alert_banner.isHidden())
        fila = self._fila_codigo(inv.tabla_inv, "S001")
        self.assertGreaterEqual(fila, 0)
        self.assertEqual(inv.tabla_inv.item(fila, 3).text(), "3 UN")
        self.assertIn("BAJO", inv.tabla_inv.item(fila, 10).text())

        self._vender(1)  # 3 - 1 = 2, sigue bajo
        inv.actualizar()
        fila = self._fila_codigo(inv.tabla_inv, "S001")
        self.assertEqual(inv.tabla_inv.item(fila, 3).text(), "2 UN")
        self.assertEqual(inv.kpi_alertas.lbl_val.text(), "1 Productos")

        self.session.expire_all()
        self.session.refresh(self.prod)
        self.assertEqual(self.prod.stock_actual, Decimal("2"))
        inv.close()

    def test_pos_e_inventario_muestran_el_mismo_stock(self):
        from ui.pos_widget import POSWidget

        self._vender(8)  # stock 2
        with patch("ui.inventario_widget.SessionLocal", self.Session):
            inv = self._inventario()
        with patch("ui.pos_widget.SessionLocal", self.Session), \
             patch("ui.pos_widget.get_session", self.fake_get_session):
            pos = POSWidget(self.user)

        fila = self._fila_codigo(inv.tabla_inv, "S001")
        stock_inv = inv.tabla_inv.item(fila, 3).text()
        stock_pos = None
        for i in range(pos._completer_model.rowCount()):
            texto = pos._completer_model.item(i).text()
            if "Coca Cola" in texto:
                stock_pos = texto
                break

        self.assertEqual(stock_inv, "2 UN")
        self.assertIsNotNone(stock_pos)
        self.assertIn("Stock: 2", stock_pos)
        inv.close()
        pos.close()


# ══════════════════════════════════════════════════════════════════════════════
#  PRUEBA 5 – COMPRAS Y AJUSTES
# ══════════════════════════════════════════════════════════════════════════════

class ComprasAjustesTest(TempDBTestCase):

    def setUp(self):
        super().setUp()
        self.cat = self.crear_categoria()
        self.prod = self.crear_producto(self.cat, "C001", "Arroz", 3, 3)
        self.user = self.crear_usuario("compras", "admin", "Usuario Compras")
        prov = Proveedor(nit="123", razon_social="Proveedor Prueba")
        self.session.add(prov)
        self.session.commit()
        self.prov = prov

    def test_compra_aumenta_stock_y_quita_alerta(self):
        self.assertTrue(self.prod.stock_bajo)
        registrar_compra(
            self.session,
            self.user,
            self.prov,
            "F-001",
            [{"producto": self.prod, "cantidad": Decimal("10"),
              "precio_unitario_bs": Decimal("5.00")}],
        )
        self.session.commit()
        self.session.refresh(self.prod)

        self.assertEqual(self.prod.stock_actual, Decimal("13"))
        self.assertNotIn(self.prod.codigo_barras,
                         [p.codigo_barras for p in productos_con_stock_bajo(self.session)])
        self.assertEqual(self.session.query(Compra).count(), 1)

    def test_ajustes_de_entrada_y_salida_se_reflejan(self):
        ajustar_stock(self.session, self.user, self.prod, Decimal("5"),
                      AjusteInventario.TIPO_ENTRADA, "Reposición")
        self.session.commit()
        self.session.refresh(self.prod)
        self.assertEqual(self.prod.stock_actual, Decimal("8"))
        self.assertFalse(self.prod.stock_bajo)

        ajustar_stock(self.session, self.user, self.prod, Decimal("6"),
                      AjusteInventario.TIPO_SALIDA, "Merma")
        self.session.commit()
        self.session.refresh(self.prod)
        self.assertEqual(self.prod.stock_actual, Decimal("2"))
        self.assertTrue(self.prod.stock_bajo)
        self.assertEqual(self.session.query(AjusteInventario).count(), 2)

    def test_inventario_refleja_cambios_externos_al_actualizar(self):
        from ui.inventario_widget import InventarioWidget

        with patch("ui.inventario_widget.SessionLocal", self.Session):
            inv = InventarioWidget(self.user)
            self.assertEqual(inv.kpi_alertas.lbl_val.text(), "1 Productos")

            # Compra registrada desde otra sesión.
            s = self.Session()
            try:
                u = s.get(Usuario, self.user.id)
                p = s.get(Producto, self.prod.id)
                prov = s.get(Proveedor, self.prov.id)
                registrar_compra(
                    s, u, prov, "F-002",
                    [{"producto": p, "cantidad": Decimal("10"),
                      "precio_unitario_bs": Decimal("5.00")}],
                )
                s.commit()
            finally:
                s.close()

            inv.actualizar()
            self.assertEqual(inv.kpi_alertas.lbl_val.text(), "0 Productos")
            inv.close()


# ══════════════════════════════════════════════════════════════════════════════
#  PRUEBA 6 – PDF Y EXCEL
# ══════════════════════════════════════════════════════════════════════════════

class ExportacionActualizadaTest(TempDBTestCase):

    def setUp(self):
        super().setUp()
        self.cat = self.crear_categoria()
        self.bajo = self.crear_producto(self.cat, "E001", "Coca Cola", 10, 3)
        self.fraccion = self.crear_producto(
            self.cat, "E002", "Arroz granel", "2.500", "3.000", unidad="KG",
        )
        self.optimo = self.crear_producto(self.cat, "E003", "Producto óptimo", 50, 5)
        self.user = self.crear_usuario("export", "cajero", "Usuario Export")
        self.session.commit()
        self.caja = self.abrir_caja(self.user)

    def test_exportaciones_reflejan_stock_actualizado(self):
        self.vender(self.user, self.bajo, 7, self.caja.id)  # queda 3 = mínimo
        bajos = productos_con_stock_bajo(self.session)
        codigos = [p.codigo_barras for p in bajos]
        self.assertIn("E001", codigos)
        self.assertIn("E002", codigos)
        self.assertNotIn("E003", codigos)

        xlsx = self.tmp_path / "bajo_stock.xlsx"
        pdf = self.tmp_path / "bajo_stock.pdf"
        exportar_productos_bajo_stock_excel(bajos, str(xlsx))
        exportar_productos_bajo_stock_pdf(bajos, str(pdf))
        self.assertTrue(xlsx.exists() and xlsx.stat().st_size > 0)
        self.assertTrue(pdf.exists() and pdf.stat().st_size > 0)

        from openpyxl import load_workbook

        wb = load_workbook(xlsx, data_only=True)
        ws = wb.active
        textos = []
        for fila in ws.iter_rows(values_only=True):
            textos.extend("" if v is None else str(v) for v in fila)
        contenido = " | ".join(textos)
        self.assertIn("E001", contenido)
        self.assertIn("E002", contenido)
        self.assertNotIn("E003", contenido)
        self.assertIn("2.5", contenido)  # cantidad fraccionaria conservada

        texto_pdf = extraer_texto_pdf(pdf.read_bytes())
        self.assertIn("E001", texto_pdf)
        self.assertIn("E002", texto_pdf)
        self.assertNotIn("E003", texto_pdf)
        self.assertIn("2.5", texto_pdf)


# ══════════════════════════════════════════════════════════════════════════════
#  PRUEBA 7 – COORDINACIÓN CENTRAL DE SEÑALES
# ══════════════════════════════════════════════════════════════════════════════

class _WidgetDummy:
    def __init__(self):
        self.llamadas = 0
        self.llamadas_caja = 0

    def actualizar(self):
        self.llamadas += 1

    def actualizar_caja_actual(self):
        self.llamadas_caja += 1


class _EmisorDummy(QWidget):
    datos_actualizados = pyqtSignal(str)


class CoordinacionSenalesTest(unittest.TestCase):

    def _main_window_sin_ui(self):
        from ui.main_window import MainWindow

        mw = MainWindow.__new__(MainWindow)
        mw._widgets = {
            "pos": _WidgetDummy(),
            "ventas": _WidgetDummy(),
            "inventario": _WidgetDummy(),
            "productos": _WidgetDummy(),
            "caja": _WidgetDummy(),
            "compras": _WidgetDummy(),
            "reportes": _WidgetDummy(),
        }
        mw._alertas_llamadas = 0
        mw._verificar_alertas = lambda: setattr(
            mw, "_alertas_llamadas", mw._alertas_llamadas + 1
        )
        return mw

    def test_venta_refresca_modulos_y_omite_emisor(self):
        from ui.main_window import MainWindow

        mw = self._main_window_sin_ui()
        MainWindow._on_datos_actualizados(mw, "pos", "venta")

        self.assertEqual(mw._widgets["ventas"].llamadas, 1)
        self.assertEqual(mw._widgets["inventario"].llamadas, 1)
        self.assertEqual(mw._widgets["productos"].llamadas, 1)
        self.assertEqual(mw._widgets["reportes"].llamadas, 1)
        self.assertEqual(mw._widgets["caja"].llamadas_caja, 1)
        self.assertEqual(mw._widgets["caja"].llamadas, 0)
        self.assertEqual(mw._widgets["pos"].llamadas, 0)  # emisor ya se refrescó
        self.assertEqual(mw._alertas_llamadas, 1)

    def test_compra_actualiza_compras_inventario_y_pos(self):
        from ui.main_window import MainWindow

        mw = self._main_window_sin_ui()
        MainWindow._on_datos_actualizados(mw, "compras", "compra")

        self.assertEqual(mw._widgets["compras"].llamadas, 0)
        self.assertEqual(mw._widgets["inventario"].llamadas, 1)
        self.assertEqual(mw._widgets["productos"].llamadas, 1)
        self.assertEqual(mw._widgets["pos"].llamadas, 1)
        self.assertEqual(mw._alertas_llamadas, 1)

    def test_conexion_de_senal_emite_y_coordina(self):
        from ui.main_window import MainWindow

        mw = self._main_window_sin_ui()
        emisor = _EmisorDummy()
        mw._widgets["pos"] = emisor
        MainWindow._conectar_senales_actualizacion(mw)

        emisor.datos_actualizados.emit("venta")

        self.assertEqual(mw._widgets["ventas"].llamadas, 1)
        self.assertEqual(mw._widgets["inventario"].llamadas, 1)
        self.assertEqual(mw._alertas_llamadas, 1)


if __name__ == "__main__":
    unittest.main(verbosity=2)
