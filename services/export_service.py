"""
services/export_service.py – Generación de reportes profesionales en PDF y Excel
para los módulos de Compras y Ventas.
"""
from __future__ import annotations

import csv
from datetime import datetime
from decimal import Decimal
from pathlib import Path
from typing import Sequence
from xml.sax.saxutils import escape as _xml_escape

# Excel (con fallback seguro en caso de no estar instalado)
try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False

# PDF (ReportLab)
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib import colors
from reportlab.lib.units import mm
from reportlab.pdfgen import canvas
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, KeepTogether
)
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle

from config import APP_NAME, CURRENCY
from models.compra import Compra
from models.producto import Producto
from models.venta import Venta


class NumberedCanvas(canvas.Canvas):
    """Canvas de ReportLab que añade numeración 'Página X de Y' y pie institucional."""
    def __init__(self, *args, **kwargs):
        super().__init__(*args, **kwargs)
        self._saved_page_states = []

    def showPage(self):
        self._saved_page_states.append(dict(self.__dict__))
        self._startPage()

    def save(self):
        num_pages = len(self._saved_page_states)
        for state in self._saved_page_states:
            self.__dict__.update(state)
            self.draw_page_decorations(num_pages)
            super().showPage()
        super().save()

    def draw_page_decorations(self, page_count: int):
        self.saveState()
        self.setFont("Helvetica", 8)
        self.setFillColor(colors.HexColor("#64748b"))

        # Línea separadora de pie
        self.setStrokeColor(colors.HexColor("#cbd5e1"))
        self.setLineWidth(0.5)
        self.line(15 * mm, 12 * mm, self._pagesize[0] - 15 * mm, 12 * mm)

        # Texto pie de página
        pie_izq = f"{APP_NAME} – Sistema de Gestión Comercial"
        pie_der = f"Página {self._pageNumber} de {page_count}"
        self.drawString(15 * mm, 8 * mm, pie_izq)
        self.drawRightString(self._pagesize[0] - 15 * mm, 8 * mm, pie_der)
        self.restoreState()


# ==============================================================================
# 1. EXPORTACIÓN DE COMPRAS (EXCEL & PDF)
# ==============================================================================

def exportar_compras_excel(compras: Sequence[Compra], filepath: str, periodo_str: str = "") -> str:
    """Genera un archivo Excel profesional (.xlsx o .csv compatible) con el listado de compras."""
    if not HAS_OPENPYXL:
        if not filepath.lower().endswith(".csv"):
            filepath = str(Path(filepath).with_suffix(".csv"))
        with open(filepath, "w", newline="", encoding="utf-8-sig") as f:
            writer = csv.writer(f)
            writer.writerow([f"{APP_NAME} - REPORTE DE COMPRAS Y ABASTECIMIENTO"])
            writer.writerow([f"Período: {periodo_str or 'Todos'}", f"Emisión: {datetime.now().strftime('%d/%m/%Y %H:%M')}"])
            writer.writerow([])
            writer.writerow(["ID", "N° Factura", "Proveedor", "NIT / CI", "Fecha y Hora", f"Total ({CURRENCY})", "Usuario", "Observaciones"])
            for c in compras:
                prov = c.proveedor.razon_social if c.proveedor else "-"
                nit = c.proveedor.nit if c.proveedor else "-"
                usr = c.usuario.username if c.usuario else "-"
                writer.writerow([c.id, c.numero_factura, prov, nit, c.fecha.strftime("%d/%m/%Y %H:%M"), f"{c.total_bs:.2f}", usr, c.observacion or ""])
            tot = sum((c.total_bs for c in compras), Decimal("0.00"))
            writer.writerow(["TOTAL:", "", "", "", "", f"{tot:.2f}", "", ""])
        return filepath

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Historial Compras"
    ws.views.sheetView[0].showGridLines = True

    # Paleta de colores
    header_fill = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")  # Azul noche
    zebra_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    kpi_fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
    total_fill = PatternFill(start_color="E2E8F0", end_color="E2E8F0", fill_type="solid")

    font_title = Font(name="Calibri", size=16, bold=True, color="0F172A")
    font_subtitle = Font(name="Calibri", size=11, italic=True, color="475569")
    font_header = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    font_data = Font(name="Calibri", size=10, color="0F172A")
    font_total = Font(name="Calibri", size=11, bold=True, color="0F172A")

    thin_border_side = Side(border_style="thin", color="CBD5E1")
    border_cell = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)
    border_total = Border(
        top=Side(border_style="thin", color="0F172A"),
        bottom=Side(border_style="double", color="0F172A"),
        left=thin_border_side,
        right=thin_border_side,
    )

    # ── Título ────────────────────────────────────────────────────────────────
    ws.merge_cells("A1:G1")
    ws["A1"] = f"{APP_NAME} – REPORTE DE COMPRAS Y ABASTECIMIENTO"
    ws["A1"].font = font_title
    ws["A1"].alignment = Alignment(vertical="center")

    ws.merge_cells("A2:G2")
    info_sub = f"Período: {periodo_str or 'Todos'}  |  Generado: {datetime.now().strftime('%d/%m/%Y %H:%M')}"
    ws["A2"] = info_sub
    ws["A2"].font = font_subtitle
    ws["A2"].alignment = Alignment(vertical="center")

    # ── Tarjetas Resumen ─────────────────────────────────────────────────────
    total_monto = sum((c.total_bs for c in compras), Decimal("0.00"))
    total_cant = len(compras)

    ws["A4"] = "Total Compras:"
    ws["A4"].font = Font(name="Calibri", size=10, bold=True, color="475569")
    ws["B4"] = f"{CURRENCY} {total_monto:,.2f}"
    ws["B4"].font = Font(name="Calibri", size=12, bold=True, color="059669")
    ws["B4"].fill = kpi_fill
    ws["B4"].alignment = Alignment(horizontal="right")

    ws["D4"] = "N° Facturas:"
    ws["D4"].font = Font(name="Calibri", size=10, bold=True, color="475569")
    ws["E4"] = total_cant
    ws["E4"].font = Font(name="Calibri", size=12, bold=True, color="0284C7")
    ws["E4"].fill = kpi_fill
    ws["E4"].alignment = Alignment(horizontal="center")

    # ── Encabezados de Tabla ─────────────────────────────────────────────────
    headers = ["ID", "N° Factura", "Proveedor", "NIT / CI", "Fecha y Hora", f"Total ({CURRENCY})", "Usuario", "Observaciones"]
    row_idx = 6

    for col_idx, h in enumerate(headers, start=1):
        c = ws.cell(row=row_idx, column=col_idx, value=h)
        c.font = font_header
        c.fill = header_fill
        c.alignment = Alignment(horizontal="center" if col_idx in (1, 2, 4, 5) else ("right" if col_idx == 6 else "left"), vertical="center")
        c.border = border_cell
    ws.row_dimensions[row_idx].height = 24

    # ── Datos ─────────────────────────────────────────────────────────────────
    for c_item in compras:
        row_idx += 1
        ws.row_dimensions[row_idx].height = 20
        nit = c_item.proveedor.nit if c_item.proveedor else "-"
        prov_nom = c_item.proveedor.razon_social if c_item.proveedor else "Sin proveedor"
        usr_nom = c_item.usuario.username if c_item.usuario else "-"

        vals = [
            c_item.id,
            c_item.numero_factura,
            prov_nom,
            nit,
            c_item.fecha.strftime("%d/%m/%Y %H:%M"),
            float(c_item.total_bs),
            usr_nom,
            c_item.observacion or "",
        ]

        is_zebra = (row_idx % 2 == 0)
        current_fill = zebra_fill if is_zebra else None

        for col_idx, val in enumerate(vals, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.font = font_data
            cell.border = border_cell
            if current_fill:
                cell.fill = current_fill

            if col_idx in (1, 2, 4, 5):
                cell.alignment = Alignment(horizontal="center", vertical="center")
            elif col_idx == 6:
                cell.alignment = Alignment(horizontal="right", vertical="center")
                cell.number_format = "#,##0.00"
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")

    # ── Fila de Totales ───────────────────────────────────────────────────────
    row_idx += 1
    ws.row_dimensions[row_idx].height = 22
    ws.merge_cells(start_row=row_idx, start_column=1, end_row=row_idx, end_column=5)
    total_lbl = ws.cell(row=row_idx, column=1, value="TOTAL GENERAL:")
    total_lbl.font = font_total
    total_lbl.alignment = Alignment(horizontal="right", vertical="center")
    total_lbl.fill = total_fill

    for ci in range(1, 6):
        ws.cell(row=row_idx, column=ci).border = border_total
        ws.cell(row=row_idx, column=ci).fill = total_fill

    tot_val = ws.cell(row=row_idx, column=6, value=float(total_monto))
    tot_val.font = Font(name="Calibri", size=11, bold=True, color="059669")
    tot_val.alignment = Alignment(horizontal="right", vertical="center")
    tot_val.number_format = "#,##0.00"
    tot_val.border = border_total
    tot_val.fill = total_fill

    ws.cell(row=row_idx, column=7).border = border_total
    ws.cell(row=row_idx, column=7).fill = total_fill
    ws.cell(row=row_idx, column=8).border = border_total
    ws.cell(row=row_idx, column=8).fill = total_fill

    # Ajuste de ancho de columnas
    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = 0
        for cell in col:
            val_str = str(cell.value or "")
            if len(val_str) > max_len and cell.row > 3:
                max_len = len(val_str)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    wb.save(filepath)
    return filepath


def exportar_compras_pdf(compras: Sequence[Compra], filepath: str, periodo_str: str = "") -> str:
    """Genera un archivo PDF profesional (.pdf) en orientación apaisada para Compras."""
    doc = SimpleDocTemplate(
        filepath,
        pagesize=landscape(A4),
        leftMargin=14 * mm,
        rightMargin=14 * mm,
        topMargin=14 * mm,
        bottomMargin=16 * mm,
    )

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "DocTitle",
        parent=styles["Normal"],
        fontName="Helvetica-Bold",
        fontSize=16,
        leading=20,
        textColor=colors.HexColor("#0f172a"),
    )
    subtitle_style = ParagraphStyle(
        "DocSubtitle",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=9,
        leading=13,
        alignment=2,  # Derecha
        textColor=colors.HexColor("#475569"),
    )
    th_style = ParagraphStyle(
        "THStyle",
        parent=styles["Normal"],
        fontName="Helvetica-Bold",
        fontSize=9,
        leading=11,
        textColor=colors.white,
        alignment=1,  # Centro
    )
    td_style = ParagraphStyle(
        "TDStyle",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=8.5,
        leading=11,
        textColor=colors.HexColor("#1e293b"),
    )
    td_center = ParagraphStyle(
        "TDCenter",
        parent=td_style,
        alignment=1,
    )
    td_right = ParagraphStyle(
        "TDRight",
        parent=td_style,
        alignment=2,
    )
    td_bold_right = ParagraphStyle(
        "TDBoldRight",
        parent=td_style,
        fontName="Helvetica-Bold",
        alignment=2,
        textColor=colors.HexColor("#059669"),
    )

    elements = []

    # ── Encabezado ────────────────────────────────────────────────────────────
    hdr_data = [
        [
            Paragraph(f"<b>{APP_NAME}</b><br/><font size=12 color='#0284c7'>Reporte de Compras y Abastecimiento</font>", title_style),
            Paragraph(f"<b>Período:</b> {periodo_str or 'Todos'}<br/><b>Emisión:</b> {datetime.now().strftime('%d/%m/%Y %H:%M')}", subtitle_style),
        ]
    ]
    t_hdr = Table(hdr_data, colWidths=[165 * mm, 104 * mm])
    t_hdr.setStyle(TableStyle([
        ('VALIGN', (0,0), (-1,-1), 'TOP'),
        ('BOTTOMPADDING', (0,0), (-1,-1), 4),
    ]))
    elements.append(t_hdr)
    elements.append(Spacer(1, 4 * mm))

    # ── Tarjetas Resumen (KPIs) ───────────────────────────────────────────────
    total_monto = sum((c.total_bs for c in compras), Decimal("0.00"))
    total_cant = len(compras)

    kpi_data = [
        [
            Paragraph("<b>TOTAL INVERTIDO EN COMPRAS</b>", ParagraphStyle("KB1", fontName="Helvetica-Bold", fontSize=8.5, textColor=colors.HexColor("#475569"))),
            Paragraph("<b>TOTAL COMPRAS REGISTRADAS</b>", ParagraphStyle("KB2", fontName="Helvetica-Bold", fontSize=8.5, textColor=colors.HexColor("#475569"))),
        ],
        [
            Paragraph(f"<font size=13 color='#059669'><b>{CURRENCY} {total_monto:,.2f}</b></font>", td_style),
            Paragraph(f"<font size=13 color='#0284c7'><b>{total_cant} compras</b></font>", td_style),
        ]
    ]
    t_kpis = Table(kpi_data, colWidths=[134 * mm, 135 * mm])
    t_kpis.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,-1), colors.HexColor("#f8fafc")),
        ('BOX', (0,0), (-1,-1), 1, colors.HexColor("#cbd5e1")),
        ('INNERGRID', (0,0), (-1,-1), 0.5, colors.HexColor("#e2e8f0")),
        ('TOPPADDING', (0,0), (-1,-1), 5),
        ('BOTTOMPADDING', (0,0), (-1,-1), 5),
        ('LEFTPADDING', (0,0), (-1,-1), 10),
    ]))
    elements.append(t_kpis)
    elements.append(Spacer(1, 5 * mm))

    # ── Tabla de Compras ──────────────────────────────────────────────────────
    col_widths = [14 * mm, 32 * mm, 68 * mm, 28 * mm, 34 * mm, 32 * mm, 30 * mm, 31 * mm]
    table_rows = [
        [
            Paragraph("<b>ID</b>", th_style),
            Paragraph("<b>N° Factura</b>", th_style),
            Paragraph("<b>Proveedor</b>", th_style),
            Paragraph("<b>NIT/CI</b>", th_style),
            Paragraph("<b>Fecha y Hora</b>", th_style),
            Paragraph(f"<b>Total ({CURRENCY})</b>", th_style),
            Paragraph("<b>Usuario</b>", th_style),
            Paragraph("<b>Observación</b>", th_style),
        ]
    ]

    for c in compras:
        prov_nom = c.proveedor.razon_social if c.proveedor else "-"
        nit = c.proveedor.nit if c.proveedor else "-"
        usr = c.usuario.username if c.usuario else "-"

        table_rows.append([
            Paragraph(str(c.id), td_center),
            Paragraph(c.numero_factura, td_center),
            Paragraph(prov_nom[:35], td_style),
            Paragraph(nit, td_center),
            Paragraph(c.fecha.strftime("%d/%m/%Y %H:%M"), td_center),
            Paragraph(f"{c.total_bs:,.2f}", td_right),
            Paragraph(usr, td_center),
            Paragraph((c.observacion or "-")[:20], td_style),
        ])

    # Fila totalizadora
    table_rows.append([
        Paragraph("<b>TOTAL:</b>", ParagraphStyle("TRTot", fontName="Helvetica-Bold", fontSize=9, alignment=2, textColor=colors.HexColor("#0f172a"))),
        "", "", "", "",
        Paragraph(f"<b>{CURRENCY} {total_monto:,.2f}</b>", td_bold_right),
        "", "",
    ])

    t_data = Table(table_rows, colWidths=col_widths, repeatRows=1)
    t_style = [
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor("#0f172a")),
        ('BOX', (0, 0), (-1, -1), 0.7, colors.HexColor("#cbd5e1")),
        ('INNERGRID', (0, 0), (-1, -2), 0.3, colors.HexColor("#e2e8f0")),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('TOPPADDING', (0, 0), (-1, -1), 4),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
        ('SPAN', (0, -1), (4, -1)),  # Unir para etiqueta TOTAL
        ('BACKGROUND', (0, -1), (-1, -1), colors.HexColor("#f1f5f9")),
        ('LINEABOVE', (0, -1), (-1, -1), 1, colors.HexColor("#0f172a")),
    ]

    # Filas alternas
    for r in range(1, len(table_rows) - 1):
        if r % 2 == 0:
            t_style.append(('BACKGROUND', (0, r), (-1, r), colors.HexColor("#f8fafc")))

    t_data.setStyle(TableStyle(t_style))
    elements.append(t_data)

    doc.build(elements, canvasmaker=NumberedCanvas)
    return filepath


# ==============================================================================
# 2. EXPORTACIÓN DE VENTAS (EXCEL & PDF)
# ==============================================================================

def exportar_ventas_excel(ventas: Sequence[Venta], filepath: str, periodo_str: str = "") -> str:
    """Genera un archivo Excel profesional (.xlsx o .csv compatible) con el listado de ventas."""
    if not HAS_OPENPYXL:
        if not filepath.lower().endswith(".csv"):
            filepath = str(Path(filepath).with_suffix(".csv"))
        with open(filepath, "w", newline="", encoding="utf-8-sig") as f:
            writer = csv.writer(f)
            writer.writerow([f"{APP_NAME} - HISTORIAL Y REPORTE DE VENTAS"])
            writer.writerow([f"Período: {periodo_str or 'Todos'}", f"Emisión: {datetime.now().strftime('%d/%m/%Y %H:%M')}"])
            writer.writerow([])
            writer.writerow(["N° Trans.", "Fecha y Hora", "Cajero", "Cliente", "NIT / CI", "Método Pago", f"Total ({CURRENCY})", "Estado"])
            for v in ventas:
                cajero = v.usuario.username if v.usuario else "-"
                cli = v.cliente.razon_social if v.cliente else "Sin cliente"
                nit = v.cliente.nit if v.cliente else "-"
                metodo = "Mixto" if v.metodo_pago == "mixto" else v.metodo_pago.capitalize()
                writer.writerow([v.numero_transaccion, v.fecha.strftime("%d/%m/%Y %H:%M"), cajero, cli, nit, metodo, f"{v.total_bs:.2f}", v.estado.upper()])
            tot = sum((v.total_bs for v in ventas if v.estado == Venta.ESTADO_COMPLETADA), Decimal("0.00"))
            writer.writerow(["TOTAL FACTURADO:", "", "", "", "", "", f"{tot:.2f}", ""])
        return filepath

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Historial Ventas"
    ws.views.sheetView[0].showGridLines = True

    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    zebra_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    kpi_fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
    total_fill = PatternFill(start_color="E2E8F0", end_color="E2E8F0", fill_type="solid")

    font_title = Font(name="Calibri", size=16, bold=True, color="0F172A")
    font_subtitle = Font(name="Calibri", size=11, italic=True, color="475569")
    font_header = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    font_data = Font(name="Calibri", size=10, color="0F172A")
    font_total = Font(name="Calibri", size=11, bold=True, color="0F172A")

    thin_border_side = Side(border_style="thin", color="CBD5E1")
    border_cell = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)
    border_total = Border(
        top=Side(border_style="thin", color="0F172A"),
        bottom=Side(border_style="double", color="0F172A"),
        left=thin_border_side,
        right=thin_border_side,
    )

    # ── Título ────────────────────────────────────────────────────────────────
    ws.merge_cells("A1:G1")
    ws["A1"] = f"{APP_NAME} – HISTORIAL Y REPORTE DE VENTAS"
    ws["A1"].font = font_title
    ws["A1"].alignment = Alignment(vertical="center")

    ws.merge_cells("A2:G2")
    info_sub = f"Período: {periodo_str or 'Todos'}  |  Generado: {datetime.now().strftime('%d/%m/%Y %H:%M')}"
    ws["A2"] = info_sub
    ws["A2"].font = font_subtitle
    ws["A2"].alignment = Alignment(vertical="center")

    # ── Tarjetas Resumen ─────────────────────────────────────────────────────
    ventas_completadas = [v for v in ventas if v.estado == Venta.ESTADO_COMPLETADA]
    total_recaudado = sum((v.total_bs for v in ventas_completadas), Decimal("0.00"))
    cant_total = len(ventas)
    cant_completadas = len(ventas_completadas)
    cant_anuladas = cant_total - cant_completadas

    ws["A4"] = "Total Facturado:"
    ws["A4"].font = Font(name="Calibri", size=10, bold=True, color="475569")
    ws["B4"] = f"{CURRENCY} {total_recaudado:,.2f}"
    ws["B4"].font = Font(name="Calibri", size=12, bold=True, color="059669")
    ws["B4"].fill = kpi_fill
    ws["B4"].alignment = Alignment(horizontal="right")

    ws["D4"] = "Ventas Completadas:"
    ws["D4"].font = Font(name="Calibri", size=10, bold=True, color="475569")
    ws["E4"] = f"{cant_completadas} operaciones"
    ws["E4"].font = Font(name="Calibri", size=11, bold=True, color="0284C7")
    ws["E4"].fill = kpi_fill
    ws["E4"].alignment = Alignment(horizontal="center")

    ws["F4"] = "Anuladas:"
    ws["F4"].font = Font(name="Calibri", size=10, bold=True, color="475569")
    ws["G4"] = f"{cant_anuladas}"
    ws["G4"].font = Font(name="Calibri", size=11, bold=True, color="EF4444")
    ws["G4"].fill = kpi_fill
    ws["G4"].alignment = Alignment(horizontal="center")

    # ── Encabezados de Tabla ─────────────────────────────────────────────────
    headers = ["N° Trans.", "Fecha y Hora", "Cajero", "Cliente", "NIT / CI", "Método Pago", f"Total ({CURRENCY})", "Estado"]
    row_idx = 6

    for col_idx, h in enumerate(headers, start=1):
        c = ws.cell(row=row_idx, column=col_idx, value=h)
        c.font = font_header
        c.fill = header_fill
        c.alignment = Alignment(horizontal="center" if col_idx in (1, 2, 5, 6, 8) else ("right" if col_idx == 7 else "left"), vertical="center")
        c.border = border_cell
    ws.row_dimensions[row_idx].height = 24

    # ── Datos ─────────────────────────────────────────────────────────────────
    for v in ventas:
        row_idx += 1
        ws.row_dimensions[row_idx].height = 20
        cli_nom = v.cliente.razon_social if v.cliente else "Sin cliente"
        cli_nit = v.cliente.nit if v.cliente else "-"
        cajero_nom = v.usuario.username if v.usuario else "-"

        if v.metodo_pago == "mixto":
            metodo_desc = "Mixto (QR + Efec)"
        elif v.metodo_pago in ("QR", "qr"):
            metodo_desc = "QR"
        else:
            metodo_desc = v.metodo_pago.capitalize()

        vals = [
            v.numero_transaccion,
            v.fecha.strftime("%d/%m/%Y %H:%M"),
            cajero_nom,
            cli_nom,
            cli_nit,
            metodo_desc,
            float(v.total_bs),
            v.estado.upper(),
        ]

        is_zebra = (row_idx % 2 == 0)
        current_fill = zebra_fill if is_zebra else None

        for col_idx, val in enumerate(vals, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.font = font_data
            cell.border = border_cell
            if current_fill:
                cell.fill = current_fill

            if col_idx == 8:
                if val == "COMPLETADA":
                    cell.font = Font(name="Calibri", size=10, bold=True, color="059669")
                else:
                    cell.font = Font(name="Calibri", size=10, bold=True, color="EF4444")

            if col_idx in (1, 2, 5, 6, 8):
                cell.alignment = Alignment(horizontal="center", vertical="center")
            elif col_idx == 7:
                cell.alignment = Alignment(horizontal="right", vertical="center")
                cell.number_format = "#,##0.00"
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")

    # ── Fila de Totales ───────────────────────────────────────────────────────
    row_idx += 1
    ws.row_dimensions[row_idx].height = 22
    ws.merge_cells(start_row=row_idx, start_column=1, end_row=row_idx, end_column=6)
    total_lbl = ws.cell(row=row_idx, column=1, value="TOTAL GENERAL EFECTIVO/VÁLIDO:")
    total_lbl.font = font_total
    total_lbl.alignment = Alignment(horizontal="right", vertical="center")
    total_lbl.fill = total_fill

    for ci in range(1, 7):
        ws.cell(row=row_idx, column=ci).border = border_total
        ws.cell(row=row_idx, column=ci).fill = total_fill

    tot_val = ws.cell(row=row_idx, column=7, value=float(total_recaudado))
    tot_val.font = Font(name="Calibri", size=11, bold=True, color="059669")
    tot_val.alignment = Alignment(horizontal="right", vertical="center")
    tot_val.number_format = "#,##0.00"
    tot_val.border = border_total
    tot_val.fill = total_fill

    ws.cell(row=row_idx, column=8).border = border_total
    ws.cell(row=row_idx, column=8).fill = total_fill

    # Ajuste de ancho de columnas
    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = 0
        for cell in col:
            val_str = str(cell.value or "")
            if len(val_str) > max_len and cell.row > 3:
                max_len = len(val_str)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    wb.save(filepath)
    return filepath


def exportar_ventas_pdf(ventas: Sequence[Venta], filepath: str, periodo_str: str = "") -> str:
    """Genera un archivo PDF profesional (.pdf) en orientación apaisada para Ventas."""
    doc = SimpleDocTemplate(
        filepath,
        pagesize=landscape(A4),
        leftMargin=14 * mm,
        rightMargin=14 * mm,
        topMargin=14 * mm,
        bottomMargin=16 * mm,
    )

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "DocTitle",
        parent=styles["Normal"],
        fontName="Helvetica-Bold",
        fontSize=16,
        leading=20,
        textColor=colors.HexColor("#0f172a"),
    )
    subtitle_style = ParagraphStyle(
        "DocSubtitle",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=9,
        leading=13,
        alignment=2,
        textColor=colors.HexColor("#475569"),
    )
    th_style = ParagraphStyle(
        "THStyle",
        parent=styles["Normal"],
        fontName="Helvetica-Bold",
        fontSize=9,
        leading=11,
        textColor=colors.white,
        alignment=1,
    )
    td_style = ParagraphStyle(
        "TDStyle",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=8.5,
        leading=11,
        textColor=colors.HexColor("#1e293b"),
    )
    td_center = ParagraphStyle(
        "TDCenter",
        parent=td_style,
        alignment=1,
    )
    td_right = ParagraphStyle(
        "TDRight",
        parent=td_style,
        alignment=2,
    )
    td_bold_right = ParagraphStyle(
        "TDBoldRight",
        parent=td_style,
        fontName="Helvetica-Bold",
        alignment=2,
        textColor=colors.HexColor("#059669"),
    )

    elements = []

    # ── Encabezado ────────────────────────────────────────────────────────────
    hdr_data = [
        [
            Paragraph(f"<b>{APP_NAME}</b><br/><font size=12 color='#0284c7'>Historial y Reporte Oficial de Ventas</font>", title_style),
            Paragraph(f"<b>Período:</b> {periodo_str or 'Todos'}<br/><b>Emisión:</b> {datetime.now().strftime('%d/%m/%Y %H:%M')}", subtitle_style),
        ]
    ]
    t_hdr = Table(hdr_data, colWidths=[165 * mm, 104 * mm])
    t_hdr.setStyle(TableStyle([
        ('VALIGN', (0,0), (-1,-1), 'TOP'),
        ('BOTTOMPADDING', (0,0), (-1,-1), 4),
    ]))
    elements.append(t_hdr)
    elements.append(Spacer(1, 4 * mm))

    # ── Tarjetas Resumen (KPIs) ───────────────────────────────────────────────
    ventas_completadas = [v for v in ventas if v.estado == Venta.ESTADO_COMPLETADA]
    total_recaudado = sum((v.total_bs for v in ventas_completadas), Decimal("0.00"))
    cant_total = len(ventas)
    cant_completadas = len(ventas_completadas)
    cant_anuladas = cant_total - cant_completadas

    kpi_data = [
        [
            Paragraph("<b>TOTAL FACTURADO</b>", ParagraphStyle("K1", fontName="Helvetica-Bold", fontSize=8.5, textColor=colors.HexColor("#475569"))),
            Paragraph("<b>OPERACIONES COMPLETADAS</b>", ParagraphStyle("K2", fontName="Helvetica-Bold", fontSize=8.5, textColor=colors.HexColor("#475569"))),
            Paragraph("<b>VENTAS ANULADAS</b>", ParagraphStyle("K3", fontName="Helvetica-Bold", fontSize=8.5, textColor=colors.HexColor("#475569"))),
        ],
        [
            Paragraph(f"<font size=13 color='#059669'><b>{CURRENCY} {total_recaudado:,.2f}</b></font>", td_style),
            Paragraph(f"<font size=13 color='#0284c7'><b>{cant_completadas} ventas</b></font>", td_style),
            Paragraph(f"<font size=13 color='#ef4444'><b>{cant_anuladas} anuladas</b></font>", td_style),
        ]
    ]
    t_kpis = Table(kpi_data, colWidths=[90 * mm, 90 * mm, 89 * mm])
    t_kpis.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,-1), colors.HexColor("#f8fafc")),
        ('BOX', (0,0), (-1,-1), 1, colors.HexColor("#cbd5e1")),
        ('INNERGRID', (0,0), (-1,-1), 0.5, colors.HexColor("#e2e8f0")),
        ('TOPPADDING', (0,0), (-1,-1), 5),
        ('BOTTOMPADDING', (0,0), (-1,-1), 5),
        ('LEFTPADDING', (0,0), (-1,-1), 10),
    ]))
    elements.append(t_kpis)
    elements.append(Spacer(1, 5 * mm))

    # ── Tabla de Ventas ───────────────────────────────────────────────────────
    col_widths = [16 * mm, 32 * mm, 30 * mm, 64 * mm, 30 * mm, 30 * mm, 32 * mm, 35 * mm]
    table_rows = [
        [
            Paragraph("<b>N° Trans.</b>", th_style),
            Paragraph("<b>Fecha y Hora</b>", th_style),
            Paragraph("<b>Cajero</b>", th_style),
            Paragraph("<b>Cliente</b>", th_style),
            Paragraph("<b>NIT/CI</b>", th_style),
            Paragraph("<b>Método</b>", th_style),
            Paragraph(f"<b>Total ({CURRENCY})</b>", th_style),
            Paragraph("<b>Estado</b>", th_style),
        ]
    ]

    for v in ventas:
        cajero_nom = v.usuario.username if v.usuario else "-"
        cli_nom = v.cliente.razon_social if v.cliente else "Sin cliente"
        cli_nit = v.cliente.nit if v.cliente else "-"
        metodo = "Mixto" if v.metodo_pago == "mixto" else v.metodo_pago.capitalize()
        estado_color = "#059669" if v.estado == Venta.ESTADO_COMPLETADA else "#ef4444"

        table_rows.append([
            Paragraph(str(v.numero_transaccion), td_center),
            Paragraph(v.fecha.strftime("%d/%m/%Y %H:%M"), td_center),
            Paragraph(cajero_nom, td_center),
            Paragraph(cli_nom[:35], td_style),
            Paragraph(cli_nit, td_center),
            Paragraph(metodo, td_center),
            Paragraph(f"{v.total_bs:,.2f}", td_right),
            Paragraph(f"<font color='{estado_color}'><b>{v.estado.upper()}</b></font>", td_center),
        ])

    table_rows.append([
        Paragraph("<b>TOTAL:</b>", ParagraphStyle("TRTotV", fontName="Helvetica-Bold", fontSize=9, alignment=2, textColor=colors.HexColor("#0f172a"))),
        "", "", "", "", "",
        Paragraph(f"<b>{CURRENCY} {total_recaudado:,.2f}</b>", td_bold_right),
        "",
    ])

    t_data = Table(table_rows, colWidths=col_widths, repeatRows=1)
    t_style = [
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor("#0f172a")),
        ('BOX', (0, 0), (-1, -1), 0.7, colors.HexColor("#cbd5e1")),
        ('INNERGRID', (0, 0), (-1, -2), 0.3, colors.HexColor("#e2e8f0")),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('TOPPADDING', (0, 0), (-1, -1), 4),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
        ('SPAN', (0, -1), (5, -1)),
        ('BACKGROUND', (0, -1), (-1, -1), colors.HexColor("#f1f5f9")),
        ('LINEABOVE', (0, -1), (-1, -1), 1, colors.HexColor("#0f172a")),
    ]

    for r in range(1, len(table_rows) - 1):
        if r % 2 == 0:
            t_style.append(('BACKGROUND', (0, r), (-1, r), colors.HexColor("#f8fafc")))

    t_data.setStyle(TableStyle(t_style))
    elements.append(t_data)

    doc.build(elements, canvasmaker=NumberedCanvas)
    return filepath


# ==============================================================================
# 3. EXPORTACIÓN DE PRODUCTOS CON BAJO STOCK (EXCEL & PDF)
# ==============================================================================

def _cantidad_texto(valor: Decimal) -> str:
    """Formatea una cantidad ya convertida desde milésimas, sin ceros sobrantes."""
    valor = Decimal(valor)
    if valor == valor.to_integral_value():
        return str(int(valor))
    return f"{valor:.3f}".rstrip("0").rstrip(".")


def _texto_pdf(valor) -> str:
    """Escapa caracteres especiales para que ReportLab no interprete XML en los textos."""
    return _xml_escape(str(valor if valor is not None else ""))


def _faltante_producto(producto: Producto) -> Decimal:
    """Cantidad faltante para alcanzar el stock mínimo (nunca negativa)."""
    return max(Decimal("0.000"), producto.stock_minimo - producto.stock_actual)


def exportar_productos_bajo_stock_excel(productos: Sequence[Producto], filepath: str) -> str:
    """Genera un Excel (.xlsx) con los productos cuyo stock actual <= stock mínimo."""
    if not HAS_OPENPYXL:
        if not filepath.lower().endswith(".csv"):
            filepath = str(Path(filepath).with_suffix(".csv"))
        with open(filepath, "w", newline="", encoding="utf-8-sig") as f:
            writer = csv.writer(f)
            writer.writerow([f"{APP_NAME} - REPORTE DE PRODUCTOS CON BAJO STOCK"])
            writer.writerow([f"Generado: {datetime.now().strftime('%d/%m/%Y %H:%M')}"])
            writer.writerow([])
            writer.writerow([
                "Código", "Producto", "Categoría", "Unidad de Medida",
                "Stock Actual", "Stock Mínimo", "Faltante para Mínimo",
            ])
            for p in productos:
                writer.writerow([
                    p.codigo_barras,
                    p.descripcion,
                    p.categoria.nombre if p.categoria else "-",
                    p.unidad_medida,
                    _cantidad_texto(p.stock_actual),
                    _cantidad_texto(p.stock_minimo),
                    _cantidad_texto(_faltante_producto(p)),
                ])
        return filepath

    generado = datetime.now()
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Bajo Stock"
    ws.views.sheetView[0].showGridLines = True

    header_fill = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")
    zebra_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    kpi_fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
    alert_fill = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")

    font_title = Font(name="Calibri", size=16, bold=True, color="0F172A")
    font_subtitle = Font(name="Calibri", size=11, italic=True, color="475569")
    font_header = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    font_data = Font(name="Calibri", size=10, color="0F172A")
    font_kpi = Font(name="Calibri", size=10, bold=True, color="475569")
    font_kpi_value = Font(name="Calibri", size=12, bold=True, color="B45309")
    font_empty = Font(name="Calibri", size=11, italic=True, color="64748B")

    thin_border_side = Side(border_style="thin", color="CBD5E1")
    border_cell = Border(
        left=thin_border_side, right=thin_border_side,
        top=thin_border_side, bottom=thin_border_side,
    )

    total_faltante = sum((_faltante_producto(p) for p in productos), Decimal("0.000"))
    agotados = sum(1 for p in productos if p.stock_actual <= 0)

    # ── Título y fecha de generación ──────────────────────────────────────────
    ws.merge_cells("A1:G1")
    ws["A1"] = f"{APP_NAME} – REPORTE DE PRODUCTOS CON BAJO STOCK"
    ws["A1"].font = font_title
    ws["A1"].alignment = Alignment(vertical="center")

    ws.merge_cells("A2:G2")
    ws["A2"] = (
        f"Generado: {generado.strftime('%d/%m/%Y %H:%M')}  |  "
        "Criterio: stock actual menor o igual al stock mínimo (productos activos)"
    )
    ws["A2"].font = font_subtitle
    ws["A2"].alignment = Alignment(vertical="center")

    # ── Resumen ───────────────────────────────────────────────────────────────
    ws["A4"] = "Productos con bajo stock:"
    ws["A4"].font = font_kpi
    ws["B4"] = len(productos)
    ws["B4"].font = font_kpi_value
    ws["B4"].fill = kpi_fill
    ws["B4"].alignment = Alignment(horizontal="center")

    ws["D4"] = "Faltante total:"
    ws["D4"].font = font_kpi
    ws["E4"] = float(total_faltante)
    ws["E4"].font = font_kpi_value
    ws["E4"].fill = kpi_fill
    ws["E4"].alignment = Alignment(horizontal="right")
    ws["E4"].number_format = "#,##0.###"

    ws["F4"] = "Agotados:"
    ws["F4"].font = font_kpi
    ws["G4"] = agotados
    ws["G4"].font = font_kpi_value
    ws["G4"].fill = kpi_fill
    ws["G4"].alignment = Alignment(horizontal="center")

    # ── Encabezados ───────────────────────────────────────────────────────────
    headers = [
        "Código", "Producto", "Categoría", "Unidad de Medida",
        "Stock Actual", "Stock Mínimo", "Faltante para Mínimo",
    ]
    row_idx = 6
    for col_idx, h in enumerate(headers, start=1):
        c = ws.cell(row=row_idx, column=col_idx, value=h)
        c.font = font_header
        c.fill = header_fill
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = border_cell
    ws.row_dimensions[row_idx].height = 30

    # ── Datos ─────────────────────────────────────────────────────────────────
    for p in productos:
        row_idx += 1
        ws.row_dimensions[row_idx].height = 20
        categoria = p.categoria.nombre if p.categoria else "-"
        vals = [
            p.codigo_barras,
            p.descripcion,
            categoria,
            p.unidad_medida,
            float(p.stock_actual),
            float(p.stock_minimo),
            float(_faltante_producto(p)),
        ]
        current_fill = zebra_fill if row_idx % 2 == 0 else None

        for col_idx, val in enumerate(vals, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.font = font_data
            cell.border = border_cell
            if current_fill:
                cell.fill = current_fill

            if col_idx in (5, 6, 7):
                cell.alignment = Alignment(horizontal="right", vertical="center")
                cell.number_format = "#,##0.###"
            elif col_idx in (1, 4):
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")

        if p.stock_actual <= 0:
            ws.cell(row=row_idx, column=5).font = Font(name="Calibri", size=10, bold=True, color="DC2626")
        else:
            ws.cell(row=row_idx, column=5).font = Font(name="Calibri", size=10, bold=True, color="B45309")
        ws.cell(row=row_idx, column=7).fill = alert_fill

    if productos:
        ws.auto_filter.ref = f"A6:G{row_idx}"
    else:
        row_idx += 1
        ws.merge_cells(start_row=row_idx, start_column=1, end_row=row_idx, end_column=7)
        empty_cell = ws.cell(
            row=row_idx, column=1,
            value="No se encontraron productos con stock actual menor o igual al stock mínimo.",
        )
        empty_cell.font = font_empty
        empty_cell.alignment = Alignment(horizontal="center", vertical="center")

    ws.freeze_panes = "A7"

    # ── Anchos de columna ─────────────────────────────────────────────────────
    max_producto = max([len(p.descripcion) for p in productos] + [len("Producto")])
    max_categoria = max(
        [len(p.categoria.nombre) if p.categoria else 1 for p in productos] + [len("Categoría")]
    )
    anchos = {
        "A": max(12, min(max([len(p.codigo_barras) for p in productos] + [len("Código")]) + 3, 24)),
        "B": max(22, min(max_producto + 3, 55)),
        "C": max(16, min(max_categoria + 3, 32)),
        "D": 18,
        "E": 14,
        "F": 14,
        "G": 21,
    }
    for col_letter, ancho in anchos.items():
        ws.column_dimensions[col_letter].width = ancho

    wb.save(filepath)
    return filepath


def exportar_productos_bajo_stock_pdf(productos: Sequence[Producto], filepath: str) -> str:
    """Genera un PDF legible e imprimible con los productos de stock bajo o agotados."""
    doc = SimpleDocTemplate(
        filepath,
        pagesize=landscape(A4),
        leftMargin=14 * mm,
        rightMargin=14 * mm,
        topMargin=14 * mm,
        bottomMargin=16 * mm,
    )

    estilos = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "DocTitleBajoStock",
        parent=estilos["Normal"],
        fontName="Helvetica-Bold",
        fontSize=16,
        leading=20,
        textColor=colors.HexColor("#0f172a"),
    )
    subtitle_style = ParagraphStyle(
        "DocSubtitleBajoStock",
        parent=estilos["Normal"],
        fontName="Helvetica",
        fontSize=9,
        leading=13,
        alignment=2,
        textColor=colors.HexColor("#475569"),
    )
    th_style = ParagraphStyle(
        "THBajoStock",
        parent=estilos["Normal"],
        fontName="Helvetica-Bold",
        fontSize=9,
        leading=11,
        textColor=colors.white,
        alignment=1,
    )
    td_style = ParagraphStyle(
        "TDBajoStock",
        parent=estilos["Normal"],
        fontName="Helvetica",
        fontSize=8.5,
        leading=11,
        textColor=colors.HexColor("#1e293b"),
    )
    td_center = ParagraphStyle("TDCenterBajoStock", parent=td_style, alignment=1)
    td_right = ParagraphStyle("TDRightBajoStock", parent=td_style, alignment=2)
    td_bold_right = ParagraphStyle(
        "TDBoldRightBajoStock",
        parent=td_style,
        fontName="Helvetica-Bold",
        alignment=2,
        textColor=colors.HexColor("#b45309"),
    )

    generado = datetime.now()
    total_faltante = sum((_faltante_producto(p) for p in productos), Decimal("0.000"))
    agotados = sum(1 for p in productos if p.stock_actual <= 0)

    elements = []

    # ── Encabezado con título y fecha ─────────────────────────────────────────
    hdr_data = [[
        Paragraph(
            f"<b>{_texto_pdf(APP_NAME)}</b><br/>"
            "<font size=12 color='#0284c7'>Reporte de Productos con Bajo Stock</font>",
            title_style,
        ),
        Paragraph(
            f"<b>Generado:</b> {generado.strftime('%d/%m/%Y %H:%M')}<br/>"
            "<b>Criterio:</b> stock actual menor o igual al stock mínimo",
            subtitle_style,
        ),
    ]]
    t_hdr = Table(hdr_data, colWidths=[165 * mm, 104 * mm])
    t_hdr.setStyle(TableStyle([
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
    ]))
    elements.append(t_hdr)
    elements.append(Spacer(1, 4 * mm))

    # ── Resumen ───────────────────────────────────────────────────────────────
    kpi_data = [
        [
            Paragraph("<b>PRODUCTOS CON BAJO STOCK</b>", ParagraphStyle("KB1", fontName="Helvetica-Bold", fontSize=8.5, textColor=colors.HexColor("#475569"))),
            Paragraph("<b>FALTANTE TOTAL (UNIDADES)</b>", ParagraphStyle("KB2", fontName="Helvetica-Bold", fontSize=8.5, textColor=colors.HexColor("#475569"))),
            Paragraph("<b>PRODUCTOS AGOTADOS</b>", ParagraphStyle("KB3", fontName="Helvetica-Bold", fontSize=8.5, textColor=colors.HexColor("#475569"))),
        ],
        [
            Paragraph(f"<font size=13 color='#b45309'><b>{len(productos)} productos</b></font>", td_style),
            Paragraph(f"<font size=13 color='#b45309'><b>{_cantidad_texto(total_faltante)}</b></font>", td_style),
            Paragraph(f"<font size=13 color='#dc2626'><b>{agotados} agotados</b></font>", td_style),
        ],
    ]
    t_kpis = Table(kpi_data, colWidths=[90 * mm, 90 * mm, 89 * mm])
    t_kpis.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#f8fafc")),
        ("BOX", (0, 0), (-1, -1), 1, colors.HexColor("#cbd5e1")),
        ("INNERGRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#e2e8f0")),
        ("TOPPADDING", (0, 0), (-1, -1), 5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
        ("LEFTPADDING", (0, 0), (-1, -1), 10),
    ]))
    elements.append(t_kpis)
    elements.append(Spacer(1, 5 * mm))

    # ── Tabla de productos ────────────────────────────────────────────────────
    col_widths = [28 * mm, 70 * mm, 43 * mm, 18 * mm, 27 * mm, 27 * mm, 30 * mm]
    table_rows = [[
        Paragraph("<b>Código</b>", th_style),
        Paragraph("<b>Producto</b>", th_style),
        Paragraph("<b>Categoría</b>", th_style),
        Paragraph("<b>Unidad</b>", th_style),
        Paragraph("<b>Stock Actual</b>", th_style),
        Paragraph("<b>Stock Mínimo</b>", th_style),
        Paragraph("<b>Faltante</b>", th_style),
    ]]

    for p in productos:
        table_rows.append([
            Paragraph(_texto_pdf(p.codigo_barras), td_center),
            Paragraph(_texto_pdf(p.descripcion), td_style),
            Paragraph(_texto_pdf(p.categoria.nombre if p.categoria else "-"), td_style),
            Paragraph(_texto_pdf(p.unidad_medida), td_center),
            Paragraph(_cantidad_texto(p.stock_actual), td_right),
            Paragraph(_cantidad_texto(p.stock_minimo), td_right),
            Paragraph(_cantidad_texto(_faltante_producto(p)), td_bold_right),
        ])

    if not productos:
        table_rows.append([
            Paragraph(
                "No se encontraron productos con stock actual menor o igual al stock mínimo.",
                td_style,
            ),
            "", "", "", "", "", "",
        ])

    t_data = Table(table_rows, colWidths=col_widths, repeatRows=1)
    t_style = [
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0f172a")),
        ("BOX", (0, 0), (-1, -1), 0.7, colors.HexColor("#cbd5e1")),
        ("INNERGRID", (0, 0), (-1, -2), 0.3, colors.HexColor("#e2e8f0")),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
    ]
    if not productos:
        t_style.append(("SPAN", (0, -1), (-1, -1)))
        t_style.append(("BACKGROUND", (0, -1), (-1, -1), colors.HexColor("#f8fafc")))

    for r in range(1, len(table_rows)):
        if productos and r % 2 == 0:
            t_style.append(("BACKGROUND", (0, r), (-1, r), colors.HexColor("#f8fafc")))

    t_data.setStyle(TableStyle(t_style))
    elements.append(t_data)

    doc.build(elements, canvasmaker=NumberedCanvas)
    return filepath
